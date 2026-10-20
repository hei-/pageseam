"""生成 Word：文字型页和混合页用 pdf2docx（保留原版式），OCR 页按版面块写入，再按原页序合并。

连续的同类页分成一组分别生成，每组是一个分节，纸张大小和方向跟随原页面。
空白页不输出。

混合页中的图片（images 参数）：
- replace（默认）：图片里识别出表格或文字时，替换成可编辑的表格和段落；照片、Logo 这类没有文字的图片保留
- keep：保留图片（外观与原件一致，图片里的文字不能编辑）
- both：保留图片，并在图片下方附上识别出的内容

跨页表格（tables 参数）：
- keep（默认）：和原件分页一致，每页一张表
- merge：合并成一张表，表头设为「重复标题行」（Word 跨页时自动重复），被分页截断的行接回上一行
"""

import copy
import io
import logging
import os
import re
from html.parser import HTMLParser

import pymupdf
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docxcompose.composer import Composer

from docparser.classify import BLANK, MIXED, PageInfo
from docparser.markdown import _PAGE_NUMBER, CAPTION_LABELS, IMAGE_LABELS, SKIP_LABELS
from docparser.ocr import Block, Result, split_items

logger = logging.getLogger(__name__)

EAST_ASIA_FONT = "宋体"
BODY_SIZE = 10.5  # 五号
TITLE_SIZES = {"doc_title": 18, "paragraph_title": 14}  # 估算不出字号时的默认值


# ---------- 文字型页：pdf2docx ----------


def convert_text_pages(pdf_path: str, page_numbers: list[int], out_path: str) -> None:
    from pdf2docx import Converter

    # pdf2docx 的日志很多，只保留警告
    logging.getLogger("pdf2docx").setLevel(logging.WARNING)
    cv = Converter(pdf_path)
    try:
        cv.convert(out_path, pages=[n - 1 for n in page_numbers], multi_processing=False)
    finally:
        cv.close()


# ---------- OCR 页：按版面块写入 ----------


class _TableParser(HTMLParser):
    """把 PP-StructureV3 输出的 HTML 表格解析成单元格网格（处理 rowspan / colspan）。"""

    def __init__(self):
        super().__init__()
        self.rows: list[list[dict]] = []
        self._cell: dict | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.rows.append([])
        elif tag in ("td", "th") and self.rows:
            a = dict(attrs)
            self._cell = {"text": "", "rowspan": int(a.get("rowspan") or 1), "colspan": int(a.get("colspan") or 1)}

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None:
            self.rows[-1].append(self._cell)
            self._cell = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell["text"] += data


def _layout_table(html: str) -> tuple[int, int, list[tuple[int, int, int, int, str]]]:
    """返回（行数，列数，[(行, 列, 跨行, 跨列, 文字)]）。"""
    parser = _TableParser()
    parser.feed(html)
    occupied: set[tuple[int, int]] = set()
    cells = []
    for r, row in enumerate(parser.rows):
        c = 0
        for cell in row:
            while (r, c) in occupied:
                c += 1
            for dr in range(cell["rowspan"]):
                for dc in range(cell["colspan"]):
                    occupied.add((r + dr, c + dc))
            cells.append((r, c, cell["rowspan"], cell["colspan"], cell["text"].strip()))
            c += cell["colspan"]
    n_rows = max((r for r, _ in occupied), default=-1) + 1
    n_cols = max((c for _, c in occupied), default=-1) + 1
    return n_rows, n_cols, cells


def _set_run_font(run, size: float | None = None, bold: bool = False) -> None:
    run.font.bold = bold or None
    if size:
        run.font.size = Pt(size)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EAST_ASIA_FONT)


def _new_document() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(BODY_SIZE)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EAST_ASIA_FONT)
    normal.paragraph_format.space_after = Pt(4)
    return doc


def _setup_section(section, size: tuple[float, float], blocks: list[Block]) -> None:
    """size 和块坐标都按转正后的版面（横着扫描的页面已经转正）。"""
    width, height = size
    section.page_width, section.page_height = Pt(width), Pt(height)
    section.orientation = WD_ORIENT.LANDSCAPE if width > height else WD_ORIENT.PORTRAIT
    # 页边距取自版面块的范围，限制在 1–3 厘米之间
    content = [b.upright for b in blocks if b.label not in SKIP_LABELS] or [pymupdf.Rect(57, 57, width - 57, height - 57)]
    left = min(r.x0 for r in content)
    right = width - max(r.x1 for r in content)
    top = min(r.y0 for r in content)
    bottom = height - max(r.y1 for r in content)
    clamp = lambda v: Cm(min(max(v / 72 * 2.54, 1.0), 3.0))  # noqa: E731
    section.left_margin, section.right_margin = clamp(left), clamp(right)
    section.top_margin, section.bottom_margin = clamp(top), clamp(min(bottom, top * 1.5))


def _alignment(block: Block, page_width: float):
    center = (block.upright.x0 + block.upright.x1) / 2
    if abs(center - page_width / 2) < page_width * 0.04 and block.upright.width < page_width * 0.7:
        return WD_ALIGN_PARAGRAPH.CENTER
    return None


def _add_table(doc: Document, html: str, width_pt: float) -> None:
    n_rows, n_cols, cells = _layout_table(html)
    if not n_rows or not n_cols:
        return
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, c, rs, cs, text in cells:
        cell = table.cell(r, c)
        if rs > 1 or cs > 1:
            cell = cell.merge(table.cell(min(r + rs, n_rows) - 1, min(c + cs, n_cols) - 1))
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        _set_run_font(paragraph.add_run(text), size=BODY_SIZE - 1)
    for column in table.columns:
        column.width = Pt(width_pt / n_cols)


def _add_image(doc: Document, page: pymupdf.Page, block: Block, angle: int) -> None:
    # 横着扫描的页面：裁剪后按转正角度旋转，和文字方向一致
    pix = page.get_pixmap(matrix=pymupdf.Matrix(200 / 72, 200 / 72).prerotate(angle), clip=block.rect)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(io.BytesIO(pix.tobytes("png")), width=Pt(block.upright.width))


def _write_blocks(doc: Document, page: pymupdf.Page, result: Result, text_width: float) -> None:
    """按版面块写入：标题、段落、表格、图片（印章等）。坐标和字号按转正后的版面。"""
    for block in result.blocks:
        text = block.content.strip()
        if block.label in SKIP_LABELS:
            continue
        if block.label == "table":
            _add_table(doc, text, min(block.upright.width, text_width))
            continue
        if block.label in IMAGE_LABELS:
            _add_image(doc, page, block, result.angle)
            continue
        if not text:
            continue
        # 字号按识别结果估算（行宽 ÷ 字数），估算不出时用默认值
        size = block.font_size or TITLE_SIZES.get(block.label) or BODY_SIZE
        # 识别结果中的换行是排版换行，不是段落结束；同一块中的多个列表项拆成多段
        for item in split_items(re.sub(r"\s*\n\s*", "", text)):
            paragraph = doc.add_paragraph()
            paragraph.alignment = _alignment(block, result.upright_size[0])
            if block.label in CAPTION_LABELS:
                paragraph.alignment = paragraph.alignment or WD_ALIGN_PARAGRAPH.LEFT
            _set_run_font(paragraph.add_run(item), size=size, bold=block.label in TITLE_SIZES)


def build_ocr_part(src: pymupdf.Document, items: list[tuple[PageInfo, Result]], out_path: str) -> None:
    doc = _new_document()
    for index, (info, result) in enumerate(items):
        if index == 0:
            section = doc.sections[0]
        else:
            section = doc.add_section(WD_SECTION.NEW_PAGE)
        _setup_section(section, result.upright_size, result.blocks)
        text_width = section.page_width.pt - section.left_margin.pt - section.right_margin.pt
        _write_blocks(doc, src[info.page - 1], result, text_width)
    doc.save(out_path)


# ---------- 混合页中的图片：替换成可编辑内容 ----------

IMAGE_MODES = ("replace", "keep", "both")
MIN_REGION_TEXT = 10  # 图片中识别出的文字少于这个字数（且没有表格）时视为照片、Logo，保留图片


def region_has_text(result: Result) -> bool:
    chars = 0
    for block in result.blocks:
        if block.label == "table":
            return True
        if block.label not in SKIP_LABELS and block.label not in IMAGE_LABELS:
            chars += len(re.sub(r"\s", "", block.content))
    return chars >= MIN_REGION_TEXT


def prepare_mixed_pages(
    src_path: str, pages: list[PageInfo], results: dict[int, list[Result]], mode: str, out_path: str
) -> dict[str, tuple[PageInfo, Result]]:
    """生成给 pdf2docx 用的副本：replace 时挖掉有文字的图片，在原位置放占位标记；both 时把标记放在图片下方。

    返回 {标记: (页, 识别结果)}，pdf2docx 转换后用 fill_placeholders 把标记换成识别出的内容。
    """
    if mode == "keep":
        return {}
    doc = pymupdf.open(src_path)
    placeholders: dict[str, tuple[PageInfo, Result]] = {}
    for info in pages:
        if info.kind != MIXED:
            continue
        page = doc[info.page - 1]
        anchors = []
        for n, (region, result) in enumerate(zip(info.ocr_regions, results.get(info.page, []))):
            if not region_has_text(result):
                continue
            rect = pymupdf.Rect(region)  # 页面显示坐标
            token = f"DPOCR{info.page}R{n}END"
            if mode == "replace":
                page.add_redact_annot(rect * page.derotation_matrix)
                anchors.append((pymupdf.Point(rect.x0, rect.y0 + 10), token))
            else:
                anchors.append((pymupdf.Point(rect.x0, rect.y1 + 12), token))
            placeholders[token] = (info, result)
        if mode == "replace" and anchors:
            # 只删除图片，保留正文和线条
            page.apply_redactions(
                images=pymupdf.PDF_REDACT_IMAGE_REMOVE,
                graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                text=pymupdf.PDF_REDACT_TEXT_NONE,
            )
        for point, token in anchors:
            page.insert_text(point * page.derotation_matrix, token, fontsize=8, fontname="helv", rotate=page.rotation)
    doc.save(out_path)
    return placeholders


def fill_placeholders(docx_path: str, src: pymupdf.Document, placeholders: dict[str, tuple[PageInfo, Result]]) -> None:
    """把 pdf2docx 输出中的占位段落换成识别出的表格和段落（写在文档末尾，再移动到占位处）。"""
    if not placeholders:
        return
    doc = Document(docx_path)
    body = doc.element.body
    section = doc.sections[0]
    text_width = section.page_width.pt - section.left_margin.pt - section.right_margin.pt
    found = set()
    for p in list(body.iter(qn("w:p"))):
        text = "".join(t.text or "" for t in p.iter(qn("w:t")))
        token = next((t for t in placeholders if t in text), None)
        if not token:
            continue
        info, result = placeholders[token]
        before = set(body)
        _write_blocks(doc, src[info.page - 1], result, min(result.upright_size[0], text_width))
        for element in [e for e in body if e not in before and e.tag != qn("w:sectPr")]:
            p.addprevious(element)
        p.getparent().remove(p)
        found.add(token)
    for token in set(placeholders) - found:
        logger.warning("placeholder %s not found in pdf2docx output", token)
    doc.save(docx_path)


# ---------- 跨页表格合并（tables=merge） ----------

TABLE_MODES = ("keep", "merge")


def _el_text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _section_break(el) -> bool:
    return el.tag == qn("w:p") and el.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None


def _filler(el) -> bool:
    """表格之间可以忽略的段落：空段落、页码。图片段落不算。"""
    if el.tag != qn("w:p") or _section_break(el) or el.find(f".//{qn('w:drawing')}") is not None:
        return False
    text = _el_text(el).strip()
    return not text or bool(_PAGE_NUMBER.match(text))


def _rows(tbl) -> list:
    return tbl.findall(qn("w:tr"))


def _cells(tr) -> list:
    return tr.findall(qn("w:tc"))


def _row_key(tr) -> list[str]:
    return [re.sub(r"\s+", "", _el_text(tc)) for tc in _cells(tr)]


def _grid_size(tbl) -> int:
    grid = tbl.find(qn("w:tblGrid"))
    return len(grid.findall(qn("w:gridCol"))) if grid is not None else 0


def _page_size(sect_pr) -> tuple | None:
    size = sect_pr.find(qn("w:pgSz")) if sect_pr is not None else None
    return (size.get(qn("w:w")), size.get(qn("w:h"))) if size is not None else None


def _next_sect_pr(children: list, start: int):
    for el in children[start:]:
        if _section_break(el):
            return el.find(f"{qn('w:pPr')}/{qn('w:sectPr')}")
        if el.tag == qn("w:sectPr"):
            return el
    return None


def _append_cell(target_tc, source_tc) -> None:
    """被截断的单元格：续接部分的第一段接到原单元格最后一段末尾，其余段落依次追加。"""
    source_ps = source_tc.findall(qn("w:p"))
    target_ps = target_tc.findall(qn("w:p"))
    if not source_ps or not _el_text(source_tc).strip():
        return
    if target_ps:
        for run in source_ps[0].findall(qn("w:r")):
            target_ps[-1].append(run)
        source_ps = source_ps[1:]
    for p in source_ps:
        target_tc.append(p)


def _merge_pair(first, second) -> None:
    rows = _rows(second)
    if rows and _row_key(rows[0]) == _row_key(_rows(first)[0]):
        second.remove(rows[0])  # 重复的表头
        rows = rows[1:]
    if rows:
        last, cont = _rows(first)[-1], rows[0]
        key = _row_key(cont)
        if len(key) == len(_cells(last)) and not key[0] and any(key[1:]):
            for target, source in zip(_cells(last), _cells(cont)):
                _append_cell(target, source)
            second.remove(cont)
            rows = rows[1:]
    for tr in rows:
        first.append(tr)
    # 表头设为「重复标题行」：Word 中表格跨页时每页自动重复
    header = _rows(first)[0]
    tr_pr = header.find(qn("w:trPr"))
    if tr_pr is None:
        tr_pr = header.makeelement(qn("w:trPr"), {})
        header.insert(0, tr_pr)
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(tr_pr.makeelement(qn("w:tblHeader"), {}))


def merge_tables(docx_path: str) -> int:
    """合并跨页表格：上一页最后一张表和下一页第一张表之间只有空段落、页码和一个分节符，
    并且列数相同或表头相同。两页纸张大小一致时同时去掉分节符。返回合并的次数。"""
    doc = Document(docx_path)
    body = doc.element.body
    merged = 0
    index = 0
    while True:
        children = list(body)
        if index >= len(children):
            break
        first = children[index]
        if first.tag != qn("w:tbl"):
            index += 1
            continue
        j, between = index + 1, []
        while j < len(children) and (_filler(children[j]) or _section_break(children[j])):
            between.append(children[j])
            j += 1
        breaks = [el for el in between if _section_break(el)]
        second = children[j] if j < len(children) else None
        if (
            len(breaks) == 1 and not _el_text(breaks[0]).strip() and second is not None and second.tag == qn("w:tbl")
            and _rows(first) and _rows(second)
            and (_grid_size(first) == _grid_size(second) or _row_key(_rows(second)[0]) == _row_key(_rows(first)[0]))
        ):
            _merge_pair(first, second)
            body.remove(second)
            same_size = _page_size(breaks[0].find(f"{qn('w:pPr')}/{qn('w:sectPr')}")) == _page_size(_next_sect_pr(children, j + 1))
            for el in between:
                if el is not breaks[0] or same_size:
                    body.remove(el)
            merged += 1
            continue  # 合并后的表格可能继续延伸到下一页
        index += 1
    if merged:
        doc.save(docx_path)
    return merged


# ---------- 合并 ----------


def _last_sect_pr(doc: Document):
    return doc.element.body.sectPr


def merge(parts: list[str], out_path: str) -> None:
    """按顺序合并多个 docx。每个部分结束处插入分节符，保留各自的纸张大小、方向和页边距。"""
    master = Document(parts[0])
    composer = Composer(master)
    for path in parts[1:]:
        # 把当前最后一节（上一部分）的节属性放进一个段落，变成分节符
        body = master.element.body
        section_end = body.makeelement(qn("w:p"), {})
        ppr = section_end.makeelement(qn("w:pPr"), {})
        ppr.append(copy.deepcopy(_last_sect_pr(master)))
        section_end.append(ppr)
        body.insert(body.index(_last_sect_pr(master)), section_end)

        part = Document(path)
        composer.append(part)
        # 文档末尾的节属性换成新部分的
        body.replace(_last_sect_pr(master), copy.deepcopy(_last_sect_pr(part)))
    composer.save(out_path)


def build(
    src_path: str,
    src: pymupdf.Document,
    pages: list[PageInfo],
    results: dict[int, list[Result]],
    work_dir: str,
    out_path: str,
    images: str = "replace",
    progress=None,
    tables: str = "keep",
) -> int:
    """pages 为本次处理的页（按顺序），results 为 {页码: [识别结果]}，images、tables 见模块说明。
    返回合并的跨页表格数（tables=keep 时为 0）。"""
    groups: list[tuple[str, list[PageInfo]]] = []
    for info in pages:
        if info.kind == BLANK:
            continue
        kind = "ocr" if info.full_ocr else "text"
        if groups and groups[-1][0] == kind:
            groups[-1][1].append(info)
        else:
            groups.append((kind, [info]))
    if not groups:
        doc = _new_document()
        doc.add_paragraph("（文档没有可转换的内容）")
        doc.save(out_path)
        return 0

    # pdf2docx 处理的副本：混合页中有文字的图片换成占位标记
    text_src = os.path.join(work_dir, "word_src.pdf")
    placeholders = prepare_mixed_pages(src_path, pages, results, images, text_src)
    if not placeholders:
        text_src = src_path

    parts = []
    for index, (kind, infos) in enumerate(groups):
        if progress:
            progress(index, len(groups))
        path = os.path.join(work_dir, f"part{index:03d}.docx")
        if kind == "text":
            convert_text_pages(text_src, [i.page for i in infos], path)
            fill_placeholders(path, src, {t: v for t, v in placeholders.items() if v[0] in infos})
        else:
            build_ocr_part(src, [(i, results[i.page][0]) for i in infos], path)
        parts.append(path)

    if len(parts) == 1:
        os.replace(parts[0], out_path)
    else:
        merge(parts, out_path)
    for path in parts + ([text_src] if text_src != src_path else []):
        if os.path.exists(path):
            os.remove(path)
    return merge_tables(out_path) if tables == "merge" else 0

