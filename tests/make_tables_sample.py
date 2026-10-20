"""生成跨页表格和 Word 样例（内容均为虚构）。

用法: python make_tables_sample.py <输出目录>

tables.pdf（5 页）：
  1-2 文字型页，物料清单跨页，第 2 页重复表头，页脚有页码
  3-4 扫描页，验收记录跨页，第 4 页没有表头
  5   文字型页，开头是另一张列数不同的表（不应和上一页合并）
office.docx：
  正文 + 跨页的长表格（设置了「重复标题行」）+ 嵌入的扫描图片（交付计划表）+ 整页扫描图片（售后服务承诺）
clauses.docx：
  条款表，其中一行的内容很长，被分页截断成两半（Word 默认「允许跨页断行」）
noheader.pdf（3 页）：
  文字型页，物料清单跨 3 页，只有第 1 页有表头（续页第一行是数据，pymupdf4llm 会把它当成表头）
noheader.docx：
  Word 长表格跨 3 页，没有设置「重复标题行」，续页没有表头
"""

import io
import os
import sys

import pymupdf

sys.path.insert(0, os.path.dirname(__file__))
from make_sample import A4, CSS, DELIVERY_HTML, SERVICE_HTML, render_scan  # noqa: E402

MATERIALS = [
    (f"WL-{1000 + i}", name, qty, price)
    for i, (name, qty, price) in enumerate(
        [("伺服电机 1.5kW", 4, 3280), ("滚珠丝杠 32×10", 2, 1460), ("直线导轨 HGH25", 6, 890), ("主轴轴承 7014", 8, 760),
         ("冷却泵 0.4kW", 2, 650), ("刀库组件 24 位", 1, 18600), ("气动夹爪", 3, 1120), ("光栅尺 500mm", 3, 2350),
         ("变频器 7.5kW", 1, 4200), ("触摸屏 10 寸", 1, 2680), ("编码器 2500 线", 5, 540), ("电磁阀 4V210", 10, 95),
         ("油水分离器", 2, 1350), ("防护罩 800mm", 4, 420), ("拖链 25×38", 6, 180), ("限位开关", 12, 36),
         ("急停按钮", 6, 48), ("三色灯", 2, 160), ("联轴器 D40", 4, 210), ("同步带 5M-450", 6, 85),
         ("同步带轮 5M-30", 6, 120), ("润滑泵 2L", 2, 980), ("油管 Φ6", 20, 12), ("接近开关 M12", 10, 45),
         ("继电器 24V", 20, 28), ("断路器 32A", 4, 95), ("接触器 25A", 4, 160), ("滤波器 20A", 2, 230),
         ("开关电源 24V", 2, 310), ("散热风扇", 6, 55), ("工作灯 LED", 2, 140), ("排屑器", 1, 5600),
         ("集中润滑分配器", 4, 75), ("气源三联件", 2, 260), ("压力表", 4, 40), ("调速阀", 6, 58)]
    )
]
ACCEPTANCE = [
    (i + 1, item, standard, result)
    for i, (item, standard, result) in enumerate(
        [("外观检查", "无磕碰、无锈蚀", "合格"), ("随机资料", "与装箱单一致", "合格"), ("定位精度 X 轴", "±0.005mm", "0.004mm"),
         ("定位精度 Y 轴", "±0.005mm", "0.003mm"), ("定位精度 Z 轴", "±0.005mm", "0.004mm"), ("重复定位精度", "±0.003mm", "0.002mm"),
         ("主轴跳动", "≤0.003mm", "0.002mm"), ("主轴温升", "≤15℃", "11℃"), ("快速移动", "48m/min", "48m/min"),
         ("换刀时间", "≤2.5s", "2.1s"), ("噪声", "≤78dB", "74dB"), ("空运转 4 小时", "无异常", "合格"),
         ("试切件尺寸", "符合图纸", "合格"), ("试切件粗糙度", "Ra≤1.6", "Ra 1.2"), ("冷却系统", "无泄漏", "合格"),
         ("润滑系统", "供油正常", "合格"), ("气动系统", "无泄漏", "合格"), ("电气安全", "接地电阻≤4Ω", "1.8Ω"),
         ("急停功能", "动作可靠", "合格"), ("连续试运行 72 小时", "无故障", "合格")]
    )
]


def material_rows(items, start):
    return "".join(
        f"<tr><td>{start + i}</td><td>{code}</td><td>{name}</td><td>{qty}</td><td>{price:,}.00</td><td>{qty * price:,}.00</td></tr>"
        for i, (code, name, qty, price) in enumerate(items)
    )


MATERIAL_HEAD = "<tr><th>序号</th><th>物料编码</th><th>名称规格</th><th>数量</th><th>单价（元）</th><th>金额（元）</th></tr>"
ACCEPT_HEAD = "<tr><th>序号</th><th>检验项目</th><th>标准要求</th><th>实测结果</th></tr>"


def accept_rows(items):
    return "".join(f"<tr><td>{n}</td><td>{item}</td><td>{std}</td><td>{res}</td></tr>" for n, item, std, res in items)


def text_page(doc, html, footer=None):
    page = doc.new_page(width=A4.width, height=A4.height)
    page.insert_htmlbox(page.rect + (50, 50, -50, -60), html, css=CSS)
    if footer:
        page.insert_text((A4.width / 2 - 10, A4.height - 30), footer, fontname="helv", fontsize=9)
    return page


def make_pdf(path):
    doc = pymupdf.open()
    # 1-2. 文字型页：物料清单跨页，第 2 页重复表头
    text_page(doc, "<h1>备品备件清单</h1><p>以下为设备首批随机备品备件，单价含税。</p>"
              f"<table>{MATERIAL_HEAD}{material_rows(MATERIALS[:26], 1)}</table>", footer="- 1 -")
    text_page(doc, f"<table>{MATERIAL_HEAD}{material_rows(MATERIALS[26:], 27)}</table>"
              "<p>以上备品备件随设备一同交付，清单外的易损件按成本价供应。</p>", footer="- 2 -")
    # 3-4. 扫描页：验收记录跨页，第 4 页没有表头
    for html in (f"<h1>设备验收记录</h1><p>设备：数控加工中心 VMC-850　编号：SN-2026-0917</p><table>{ACCEPT_HEAD}{accept_rows(ACCEPTANCE[:12])}</table>",
                 f"<table>{accept_rows(ACCEPTANCE[12:])}</table>"):
        page = doc.new_page(width=A4.width, height=A4.height)
        page.insert_image(page.rect, stream=render_scan(html, seal=False, angle=0.3))
    # 5. 另一张表（列数不同），紧接在上一页的表格之后，不应合并
    text_page(doc, "<table><tr><th>联系人</th><th>电话</th></tr><tr><td>王工（售后）</td><td>400-820-6688</td></tr></table>"
              "<p>验收中发现的问题请联系售后服务热线。</p>")
    doc.save(path, garbage=4, deflate=True)


def make_docx(path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    d = Document()
    d.add_heading("设备维保服务合同（Word 样例）", 1)
    d.add_paragraph("甲方委托乙方对下列设备提供年度维保服务，服务期自 2026 年 11 月 1 日至 2027 年 10 月 31 日。")
    rows = [("序号", "物料编码", "名称规格", "数量", "单价（元）", "金额（元）")] + [
        (str(i + 1), code, name, str(qty), f"{price:,}.00", f"{qty * price:,}.00") for i, (code, name, qty, price) in enumerate(MATERIALS)
    ]
    table = d.add_table(rows=len(rows), cols=6)
    table.style = "Table Grid"
    for r, values in enumerate(rows):
        for c, v in enumerate(values):
            table.cell(r, c).text = v
    header = OxmlElement("w:tblHeader")  # 「重复标题行」：表格跨页时每页都显示表头
    header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    d.add_paragraph("以上备品备件随设备一同交付。")
    d.add_paragraph("附件四 交付计划（扫描件）：")
    landscape = pymupdf.paper_rect("a4-l")
    d.add_picture(io.BytesIO(render_scan(DELIVERY_HTML, rect=landscape, seal=False, angle=0.2)), width=Inches(6.2))
    d.add_page_break()
    d.add_picture(io.BytesIO(render_scan(SERVICE_HTML, seal=False, angle=-0.4)), width=Inches(6.2))
    d.save(path)


CLAUSES = [("交货期限", "乙方应于合同签订后 45 个工作日内交货。"), ("交货地点", "甲方一号厂房。"),
           ("运输费用", "由乙方承担，含装卸费用。"), ("包装要求", "木箱包装，适合长途运输和多次装卸。"),
           ("开箱检验", "到货后 3 日内双方共同开箱检验。"), ("安装调试", "乙方负责，甲方提供水电和吊装条件。"),
           ("技术培训", "现场培训不少于 3 天。"), ("验收标准", "按附件五执行。"), ("付款方式", "30% 预付款，60% 验收款，10% 质保金。"),
           ("质保期限", "验收合格之日起 24 个月。"), ("售后响应", "接到通知后 24 小时内响应，72 小时内到场。"),
           ("备件供应", "常用备品备件保证供应 8 年。"), ("知识产权", "乙方保证所供设备不侵犯第三方知识产权。"),
           ("保密义务", "双方对合同内容和技术资料负有保密义务。")]
# 长条款：放在表格靠后的位置，被分页截断，开头和结尾分别在两页上
LONG_CLAUSE = ("违约责任", "".join(f"第 {i} 款：乙方逾期交货的，每逾期一日按合同总额的千分之零点五支付违约金。" for i in range(1, 16))
               + "以上违约金合计不超过合同总额的百分之十。")


def make_clauses_docx(path):
    from docx import Document

    d = Document()
    d.add_heading("合同主要条款（Word 样例）", 1)
    rows = [("序号", "条款", "内容")] + [(str(i + 1), k, v) for i, (k, v) in enumerate(CLAUSES)]
    rows += [(str(len(CLAUSES) + 1), *LONG_CLAUSE), (str(len(CLAUSES) + 2), "争议解决", "协商不成的，向甲方所在地人民法院起诉。")]
    table = d.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    for r, values in enumerate(rows):
        for c, v in enumerate(values):
            table.cell(r, c).text = v
    d.save(path)


# 没有表头的续页：90 行物料（36 种物料循环，编号连续）
LONG_ITEMS = [(f"WL-{2000 + i}", MATERIALS[i % len(MATERIALS)][1], MATERIALS[i % len(MATERIALS)][2], MATERIALS[i % len(MATERIALS)][3])
              for i in range(90)]


def make_noheader_pdf(path):
    doc = pymupdf.open()
    text_page(doc, "<h1>年度备件采购清单</h1><p>以下为全年备件采购计划，续页不重复表头。</p>"
              f"<table>{MATERIAL_HEAD}{material_rows(LONG_ITEMS[:30], 1)}</table>", footer="- 1 -")
    text_page(doc, f"<table>{material_rows(LONG_ITEMS[30:65], 31)}</table>", footer="- 2 -")
    text_page(doc, f"<table>{material_rows(LONG_ITEMS[65:], 66)}</table><p>采购清单结束。</p>", footer="- 3 -")
    doc.save(path, garbage=4, deflate=True)


def make_noheader_docx(path):
    from docx import Document

    d = Document()
    d.add_heading("年度备件采购清单（Word 样例，续页无表头）", 1)
    rows = [("序号", "物料编码", "名称规格", "数量", "单价（元）", "金额（元）")] + [
        (str(i + 1), code, name, str(qty), f"{price:,}.00", f"{qty * price:,}.00") for i, (code, name, qty, price) in enumerate(LONG_ITEMS)
    ]
    table = d.add_table(rows=len(rows), cols=6)
    table.style = "Table Grid"
    for r, values in enumerate(rows):
        for c, v in enumerate(values):
            table.cell(r, c).text = v
    d.add_paragraph("采购清单结束。")
    d.save(path)


def main(out_dir):
    os.makedirs(out_dir, exist_ok=True)
    make_pdf(os.path.join(out_dir, "tables.pdf"))
    make_docx(os.path.join(out_dir, "office.docx"))
    make_clauses_docx(os.path.join(out_dir, "clauses.docx"))
    make_noheader_pdf(os.path.join(out_dir, "noheader.pdf"))
    make_noheader_docx(os.path.join(out_dir, "noheader.docx"))
    print(f"saved tables.pdf, office.docx, clauses.docx, noheader.pdf, noheader.docx to {out_dir}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else ".")
